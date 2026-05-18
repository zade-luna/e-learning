import logging
from pathlib import Path
from dataclasses import dataclass, field
from typing import Optional

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

logger = logging.getLogger(__name__)

# If a PDF page has fewer than this many characters it is treated as scanned
SCANNED_PAGE_THRESHOLD = 50
# If more than this fraction of pages are scanned, reject the whole file
SCANNED_RATIO_THRESHOLD = 0.5


@dataclass
class SlideContent:
    slide_number: int
    title: Optional[str]
    body: str


@dataclass
class ExtractionResult:
    slides: list[SlideContent] = field(default_factory=list)
    is_scanned: bool = False
    error: Optional[str] = None

    @property
    def has_content(self) -> bool:
        return any(s.body.strip() or (s.title and s.title.strip()) for s in self.slides)


# ──────────────────────────────────────────────
# PDF
# ──────────────────────────────────────────────

def extract_from_pdf(path: Path) -> ExtractionResult:
    doc = fitz.open(str(path))
    slides: list[SlideContent] = []
    scanned_count = 0

    for page_num, page in enumerate(doc, 1):
        text = page.get_text("text").strip()

        if len(text) < SCANNED_PAGE_THRESHOLD:
            scanned_count += 1
            slides.append(SlideContent(
                slide_number=page_num,
                title=f"Page {page_num}",
                body=""
            ))
            continue

        lines = [l.strip() for l in text.split("\n") if l.strip()]
        title = lines[0] if lines else f"Page {page_num}"
        body = " ".join(lines[1:]) if len(lines) > 1 else text

        slides.append(SlideContent(
            slide_number=page_num,
            title=title,
            body=body
        ))

    doc.close()

    total = len(slides)
    if total > 0 and (scanned_count / total) > SCANNED_RATIO_THRESHOLD:
        return ExtractionResult(
            slides=slides,
            is_scanned=True,
            error=(
                f"This PDF appears to be a scanned document "
                f"({scanned_count}/{total} pages have no text layer). "
                "OCR support is not yet available."
            )
        )

    return ExtractionResult(slides=slides)


# ──────────────────────────────────────────────
# PPTX
# ──────────────────────────────────────────────

_TITLE_PLACEHOLDER_TYPES = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
}


def extract_from_pptx(path: Path) -> ExtractionResult:
    prs = Presentation(str(path))
    slides: list[SlideContent] = []

    for slide_num, slide in enumerate(prs.slides, 1):
        title_text: Optional[str] = None
        body_parts: list[str] = []

        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue

            # Identify title placeholder
            try:
                if (
                    shape.placeholder_format is not None
                    and shape.placeholder_format.type in _TITLE_PLACEHOLDER_TYPES
                ):
                    title_text = shape.text_frame.text.strip() or None
                    continue
            except Exception:
                pass

            # Everything else is body content
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    body_parts.append(text)

        # Skip completely empty slides (e.g. divider slides with only an image)
        if not title_text and not body_parts:
            continue

        slides.append(SlideContent(
            slide_number=slide_num,
            title=title_text or f"Slide {slide_num}",
            body=" ".join(body_parts)
        ))

    return ExtractionResult(slides=slides)


# ──────────────────────────────────────────────
# Router
# ──────────────────────────────────────────────

def extract_text(path: Path) -> ExtractionResult:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return extract_from_pdf(path)
    elif suffix == ".pptx":
        return extract_from_pptx(path)
    elif suffix == ".ppt":
        # python-pptx does NOT support old binary .ppt — only .pptx (Office Open XML)
        return ExtractionResult(
            error=(
                "Old-format .ppt files are not supported. "
                "Please ask the teacher to re-save as .pptx in PowerPoint: "
                "File -> Save As -> PowerPoint Presentation (.pptx)."
            )
        )
    else:
        return ExtractionResult(
            error=f"Unsupported file type: '{suffix}'. Only .pdf and .pptx are supported."
        )
